import os
import shutil
from typing import List, Tuple, Optional
import streamlit as st
import openai
from openai import OpenAI, AzureOpenAI
import chromadb
from chromadb.config import Settings
import PyPDF2
from docx import Document
import tiktoken
from langchain.text_splitter import RecursiveCharacterTextSplitter
import json
import logging

# Setup logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class ContractRAGEngine:
    def __init__(self, api_key: str = None, use_azure: bool = False):
        """Initialize RAG engine with explicit API key parameter"""
        self.use_azure = use_azure
        self.api_key = api_key  # Store the API key explicitly
        self.setup_directories()
        self.setup_openai_client()
        self.setup_vector_db()
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=1000,
            chunk_overlap=200,
            separators=["\n\n", "\n", ".", "!", "?", ",", " ", ""]
        )
    
    def setup_directories(self):
        """Create necessary directories"""
        data_dir = os.getenv("DATA_VOLUME_PATH", "data")
        embeddings_dir = os.getenv("EMBEDDINGS_VOLUME_PATH", "embeddings")
        
        os.makedirs(data_dir, exist_ok=True)
        os.makedirs(embeddings_dir, exist_ok=True)
        
        self.data_dir = data_dir
        self.embeddings_dir = embeddings_dir
    
    def _is_placeholder_key(self, key: str) -> bool:
        """Check if key is a placeholder value"""
        if not key or not key.strip():
            return True
        
        key = key.strip()
        
        # Check for OpenAI and Azure placeholders
        common_placeholders = [
            "your-openai-api-key-here",
            "sk-your-openai-api-key-here",
            "your-azure-openai-key-here"
        ]
        
        return key.lower() in [p.lower() for p in common_placeholders]
    
    def _get_valid_api_key(self) -> str:
        """Get valid API key - ONLY for Personal mode"""
        # Priority 1: Explicitly passed API key
        if self.api_key and not self._is_placeholder_key(self.api_key):
            return self.api_key
        
        # Priority 2: Session state (user entered in UI)
        try:
            session_key = st.session_state.get("temp_openai_key")
            if session_key and not self._is_placeholder_key(session_key):
                return session_key
        except:
            pass
        
        # Priority 3: Environment variable (only if not placeholder)
        env_key = os.getenv("OPENAI_API_KEY")
        if env_key and not self._is_placeholder_key(env_key):
            return env_key
        
        # No valid key found
        raise ValueError("No valid Personal OpenAI API key found")
    
    def setup_openai_client(self):
        """Setup OpenAI client with improved key handling"""
        try:
            if self.use_azure:
                # Azure OpenAI setup
                azure_key = os.getenv("AZURE_OPENAI_KEY")
                azure_endpoint = os.getenv("AZURE_OPENAI_ENDPOINT")
                
                # Check for missing credentials
                if not azure_key or not azure_endpoint:
                    raise ValueError("Azure OpenAI credentials not found in environment variables")
                
                # Check for placeholder values - BOTH key and endpoint
                if (self._is_placeholder_key(azure_key) or 
                    "your-resource" in azure_endpoint.lower()):
                    raise ValueError("Azure OpenAI credentials are placeholder values")
                
                self.client = AzureOpenAI(
                    api_key=azure_key,
                    api_version="2023-12-01-preview",
                    azure_endpoint=azure_endpoint
                )
                self.embedding_model = os.getenv("AZURE_EMBEDDING_DEPLOYMENT", "text-embedding-ada-002")
                self.chat_model = os.getenv("AZURE_CHAT_DEPLOYMENT", "gpt-35-turbo")
                logger.info("✅ Azure OpenAI client initialized successfully")
            else:
                # Personal OpenAI setup
                api_key = self._get_valid_api_key()
                
                if not api_key.startswith("sk-") or len(api_key) < 20:
                    raise ValueError("Invalid OpenAI API key format")
                
                self.client = OpenAI(api_key=api_key)
                self.embedding_model = os.getenv("OPENAI_EMBEDDING_MODEL", "text-embedding-3-small")
                self.chat_model = os.getenv("OPENAI_CHAT_MODEL", "gpt-3.5-turbo")
                logger.info("✅ Personal OpenAI client initialized successfully")
                
        except Exception as e:
            logger.error(f"❌ OpenAI setup failed: {e}")
            raise e
    
    def setup_vector_db(self):
        """Setup ChromaDB vector database"""
        try:
            self.chroma_client = chromadb.PersistentClient(
                path=self.embeddings_dir,
                settings=Settings(anonymized_telemetry=False)
            )
            self.collection = self.chroma_client.get_or_create_collection(
                name="contracts",
                metadata={"hnsw:space": "cosine"}
            )
            logger.info(f"✅ ChromaDB initialized at {self.embeddings_dir}")
        except Exception as e:
            logger.error(f"❌ ChromaDB setup failed: {e}")
            raise e
    
    def test_api_connection(self) -> Tuple[bool, str]:
        """Test API connection with a simple embedding call"""
        try:
            test_response = self.client.embeddings.create(
                model=self.embedding_model,
                input="test connection"
            )
            if test_response and test_response.data:
                return True, "✅ API connection successful"
            else:
                return False, "❌ API returned empty response"
        except Exception as e:
            return False, f"❌ API connection failed: {str(e)}"
    
    def get_hello_message(self, name: str) -> str:
        """Return greeting message"""
        return f"Hello {name}! Welcome to the Contract Analyzer Agent."
    
    def is_valid_file(self, uploaded_file) -> Tuple[bool, str]:
        """Validate uploaded file format and size"""
        valid_extensions = ['.pdf', '.docx', '.doc']
        max_size = 10 * 1024 * 1024  # 10MB limit
        
        file_extension = os.path.splitext(uploaded_file.name)[1].lower()
        
        if file_extension not in valid_extensions:
            return False, f"Invalid file format. Only PDF and DOC/DOCX files are allowed."
        
        if uploaded_file.size > max_size:
            return False, f"File too large. Maximum size is 10MB."
        
        return True, "File is valid"
    
    def save_uploaded_file(self, uploaded_file, replace_existing: bool = True) -> Tuple[bool, str]:
        """Save uploaded file with option to replace existing"""
        try:
            is_valid, message = self.is_valid_file(uploaded_file)
            if not is_valid:
                return False, message
            
            file_path = os.path.join(self.data_dir, uploaded_file.name)
            
            # Handle existing file
            if os.path.exists(file_path):
                if replace_existing:
                    logger.info(f"Replacing existing file: {uploaded_file.name}")
                    # Remove from vector database first
                    self._remove_document_from_db(uploaded_file.name)
                    os.remove(file_path)
                else:
                    return False, f"File '{uploaded_file.name}' already exists."
            
            # Save new file
            with open(file_path, "wb") as f:
                f.write(uploaded_file.getbuffer())
            
            logger.info(f"✅ File saved: {file_path}")
            return True, f"File saved successfully: {uploaded_file.name}"
            
        except Exception as e:
            logger.error(f"❌ File save error: {e}")
            return False, f"Error saving file: {str(e)}"
    
    def _remove_document_from_db(self, filename: str):
        """Remove document chunks from vector database"""
        try:
            # Get all chunk IDs for this document
            all_data = self.collection.get()
            ids_to_delete = []
            
            for i, metadata in enumerate(all_data['metadatas']):
                if metadata.get('source') == filename:
                    ids_to_delete.append(all_data['ids'][i])
            
            if ids_to_delete:
                self.collection.delete(ids=ids_to_delete)
                logger.info(f"Removed {len(ids_to_delete)} chunks for {filename}")
                
        except Exception as e:
            logger.error(f"Error removing document from DB: {e}")
    
    def extract_text_from_pdf(self, file_path: str) -> str:
        """Extract text from PDF file"""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                text_parts = []
                for page_num, page in enumerate(pdf_reader.pages):
                    try:
                        page_text = page.extract_text()
                        if page_text.strip():
                            text_parts.append(page_text.strip())
                    except Exception as page_error:
                        logger.warning(f"Error extracting page {page_num} from {file_path}: {page_error}")
                        continue
                return "\n\n".join(text_parts)
        except Exception as e:
            logger.error(f"PDF extraction error for {file_path}: {e}")
            return ""
    
    def extract_text_from_docx(self, file_path: str) -> str:
        """Extract text from DOCX file"""
        try:
            doc = Document(file_path)
            text_parts = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text.strip())
            return "\n".join(text_parts)
        except Exception as e:
            logger.error(f"DOCX extraction error for {file_path}: {e}")
            return ""
    
    def extract_text_from_file(self, file_path: str) -> str:
        """Extract text based on file extension"""
        file_extension = os.path.splitext(file_path)[1].lower()
        
        if file_extension == '.pdf':
            return self.extract_text_from_pdf(file_path)
        elif file_extension in ['.docx', '.doc']:
            return self.extract_text_from_docx(file_path)
        else:
            return ""
    
    def get_embedding(self, text: str) -> List[float]:
        """Get embedding for text using OpenAI"""
        try:
            response = self.client.embeddings.create(
                model=self.embedding_model,
                input=text
            )
            return response.data[0].embedding
        except Exception as e:
            logger.error(f"❌ Embedding error: {e}")
            return []
    
    def process_and_store_document(self, file_path: str) -> Tuple[bool, str]:
        """Process document and store in vector database with detailed feedback"""
        try:
            filename = os.path.basename(file_path)
            logger.info(f"🔄 Processing document: {filename}")
            
            # Extract text
            text = self.extract_text_from_file(file_path)
            if not text.strip():
                return False, f"No text could be extracted from {filename}"
            
            logger.info(f"✅ Extracted {len(text)} characters from {filename}")
            
            # Split into chunks
            chunks = self.text_splitter.split_text(text)
            if not chunks:
                return False, f"No text chunks created from {filename}"
            
            logger.info(f"📄 Created {len(chunks)} chunks from {filename}")
            
            # Store in ChromaDB
            stored_chunks = 0
            failed_chunks = 0
            
            for i, chunk in enumerate(chunks):
                if not chunk.strip():
                    continue
                
                chunk_id = f"{filename}_chunk_{i}"
                
                # Get embedding
                embedding = self.get_embedding(chunk)
                if not embedding:
                    failed_chunks += 1
                    logger.warning(f"⚠️ Failed to get embedding for chunk {i} of {filename}")
                    continue
                
                # Store in vector database
                try:
                    self.collection.add(
                        documents=[chunk],
                        embeddings=[embedding],
                        metadatas=[{
                            "source": filename,
                            "chunk_id": i,
                            "file_path": file_path
                        }],
                        ids=[chunk_id]
                    )
                    stored_chunks += 1
                except Exception as chunk_error:
                    failed_chunks += 1
                    logger.error(f"❌ Error storing chunk {i} of {filename}: {chunk_error}")
                    continue
            
            if stored_chunks > 0:
                success_msg = f"✅ Successfully processed {filename}: {stored_chunks}/{len(chunks)} chunks stored"
                if failed_chunks > 0:
                    success_msg += f" ({failed_chunks} chunks failed)"
                logger.info(success_msg)
                return True, success_msg
            else:
                error_msg = f"❌ Failed to store any chunks from {filename}"
                logger.error(error_msg)
                return False, error_msg
                
        except Exception as e:
            error_msg = f"❌ Document processing error for {file_path}: {e}"
            logger.error(error_msg)
            return False, error_msg
    
    def get_relevant_chunks(self, query: str, n_results: int = 5) -> List[dict]:
        """Retrieve relevant chunks for query"""
        try:
            query_embedding = self.get_embedding(query)
            if not query_embedding:
                return []
            
            results = self.collection.query(
                query_embeddings=[query_embedding],
                n_results=n_results
            )
            
            chunks = []
            for i in range(len(results['documents'][0])):
                chunks.append({
                    'content': results['documents'][0][i],
                    'metadata': results['metadatas'][0][i],
                    'distance': results['distances'][0][i] if 'distances' in results else 0
                })
            
            return chunks
            
        except Exception as e:
            logger.error(f"❌ Retrieval error: {e}")
            return []
    
    def generate_response(self, query: str, context_chunks: List[dict]) -> str:
        """Generate response using OpenAI chat completion"""
        try:
            context = "\n\n".join([
                f"From {chunk['metadata']['source']}: {chunk['content']}"
                for chunk in context_chunks
            ])
            
            # Updated system prompt - no citations/sources in response
            system_prompt = """You are a Contract Analyzer Agent helping a procurement team analyze IT contracts and SOWs.

Your role:
- Analyze contract documents to answer procurement-related questions
- Provide accurate, specific information based on the provided context
- Focus on contract values, vendors, technologies, expiration dates, and legal requirements
- Never fabricate information - only use the provided context
- If information is not available in the provided context, clearly state "The requested information is not provided in the available contract documents"

Response format:
- Give direct, professional answers
- Do NOT include document citations, source references, or bracketed file names in your response
- If information is not available in the provided context, clearly state this
- For numerical questions (counts, values), be precise
- For expiration dates, only mention if explicit expiration dates are found in the contracts"""

            user_prompt = f"""Based on the following contract documents, please answer this question:

{query}

Context from contracts:
{context}

Please provide a comprehensive answer without including any citations or source references."""
            
            response = self.client.chat.completions.create(
                model=self.chat_model,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt}
                ],
                temperature=0.1,
                max_tokens=1000
            )
            
            return response.choices[0].message.content
            
        except Exception as e:
            logger.error(f"❌ Response generation error: {e}")
            return f"Error generating response: {str(e)}"
    
    def query_contracts(self, query: str) -> Tuple[str, List[str]]:
        """Main query function - retrieve and generate response"""
        try:
            chunks = self.get_relevant_chunks(query, n_results=5)
            
            if not chunks:
                return "No relevant information found in the uploaded contracts.", []
            
            response = self.generate_response(query, chunks)
            sources = list(set([chunk['metadata']['source'] for chunk in chunks]))
            
            return response, sources
            
        except Exception as e:
            logger.error(f"❌ Query error: {e}")
            return f"Error processing query: {str(e)}", []
    
    def get_uploaded_files(self) -> List[str]:
        """Get list of uploaded files"""
        try:
            if not os.path.exists(self.data_dir):
                return []
            return [f for f in os.listdir(self.data_dir) if f.endswith(('.pdf', '.docx', '.doc'))]
        except Exception as e:
            logger.error(f"❌ Error getting file list: {e}")
            return []
    
    def get_collection_stats(self) -> dict:
        """Get statistics about the vector database"""
        try:
            count = self.collection.count()
            files = self.get_uploaded_files()
            return {
                "total_chunks": count,
                "total_files": len(files),
                "files": files
            }
        except Exception as e:
            logger.error(f"❌ Error getting stats: {e}")
            return {"total_chunks": 0, "total_files": 0, "files": []}
    
    def clear_all_data(self):
        """Clear all uploaded files and embeddings"""
        try:
            # Clear vector database
            self.chroma_client.delete_collection("contracts")
            self.collection = self.chroma_client.get_or_create_collection(
                name="contracts",
                metadata={"hnsw:space": "cosine"}
            )
            
            # Clear uploaded files
            if os.path.exists(self.data_dir):
                shutil.rmtree(self.data_dir)
            os.makedirs(self.data_dir, exist_ok=True)
            
            logger.info("✅ All data cleared successfully")
            return True
        except Exception as e:
            logger.error(f"❌ Error clearing data: {e}")
            return False